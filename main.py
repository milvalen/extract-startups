import os
import requests
from docx import Document
from docx.table import Table
from openpyxl import Workbook

PATH = './data/'
OUTPUT_PATH = './output/res.xlsx'
SEPARATOR = 'ПАСПОРТ СТАРТАП-ПРОЕКТА'
BREAK = '-textbreak-'


def check_url(url):
    try:
        response = requests.get(url, timeout=10, allow_redirects=False)
        if response.status_code == 200:
            return True
        else:
            return False
    except requests.exceptions.RequestException:
        return False


def extract_startups(text):
    links_num = 0
    startups = []

    if len(text.split(SEPARATOR)) < 2:
        text = text.replace('Тема стартап-проекта*', SEPARATOR)
    print(len(text.split(SEPARATOR)))
    for passport in text.split(SEPARATOR):
        startup = {'Название стартап-проекта': '', 'Ссылка': ''}

        passport_parts = passport.split(BREAK)
        for part in passport_parts:
            joined_part = ''.join(part)
            refined_part = (
                joined_part
                .replace(' ', '')
                .replace(',', '.')
                .replace('\xad', '')
                .replace('\\', 'l')
                .replace('е', 'e')
                .replace('Т', 'T')
                .replace('у', 'y')
                .replace('о', 'o')
                .replace('р', 'p')
                .replace('А', 'A')
                .replace('а', 'a')
                .replace('Н', 'H')
                .replace('О', 'O')
                .replace('Р', 'P')
                .replace('К', 'K')
                .replace('Х', 'X')
                .replace('х', 'x')
                .replace('В', 'B')
                .replace('М', 'M')
                .replace('г', 'r')
                .replace('Р', 'P')
                .replace('С', 'C')
                .replace('с', 'c')
                .replace('Е', 'E')
                .replace('ь', 'b')
                .replace('т', 'T')
                .replace('п', 'n')
                .replace('У', 'Y')
                .replace('З', '3')
                .replace('з', '3')
                .replace('ш', 'w')
                .replace('Ш', 'W')
                .replace('Ы', 'bl')
                .replace('ы', 'bl')
                .replace('ц', 'u')
                .replace('Ц', 'U')
                .replace('Ь', 'b')
                .split('\t'))[0]
            if 'https://pt.2035' in refined_part and not startup['Ссылка']:
                startup['Ссылка'] = refined_part
                next_line = passport_parts[passport_parts.index(part) + 1]

                if refined_part.endswith('-') \
                        and next_line \
                        and not (next_line[0].isdigit() or next_line.startswith('Наименование')):
                    startup['Ссылка'] += next_line
            elif 'Название стартап-проекта*' in joined_part:
                part_index = passport_parts.index(part)
                name_shift = 1

                while not passport_parts[part_index + name_shift]:
                    name_shift += 1

                startup_name = passport_parts[part_index + name_shift].replace('\t', '')
                if not startup_name.endswith('*'):
                    startup['Название стартап-проекта'] = startup_name

        if startup['Название стартап-проекта']:
            startups.append(startup)
            print(startup)
            if startup['Ссылка']:
                links_num += 1
    print('\nПроверяются ссылки...')

    for startup in startups:
        if startup['Ссылка'] and not check_url(startup['Ссылка']):
            print('Нерабочая ссылка -', startup)

    startups_num = len(startups)
    print('\n{} строк, {} из них без ссылок'.format(startups_num, startups_num - links_num))

    return startups


def iter_block_items(doc):
    for block in doc.element.body:
        yield block


def extract_text_in_order(docs):
    full_text = []
    for doc in docs:
        for block in iter_block_items(doc):
            if type(block).__name__ == 'CT_Tbl':
                table_text = []
                table = Table(block, doc)
                for row in table.rows:
                    prior_tc = None
                    for cell in row.cells:
                        if cell._tc == prior_tc:
                            continue
                        prior_tc = cell._tc
                        for paragraph in cell.paragraphs:
                            table_text.append(paragraph.text + BREAK)
                            # print(paragraph.text)

                full_text.append(''.join(table_text))
            elif block.text:
                # print(block.text)
                full_text.append(block.text.replace(
                    'Паспорт стартап-проекта', SEPARATOR) + BREAK)

    return extract_startups(''.join(full_text))


def process_dir(path):
    docs = []

    for root, _, files in os.walk(path):
        for file in files:
            if file.endswith('.docx'):
                print(file)
                word_path = os.path.join(root, file)
                if isinstance(word_path, str):
                    docs.append(Document(word_path))

    return extract_text_in_order(docs)


def save_to_excel(data, excel_filename):
    wb = Workbook()
    ws = wb.active

    headers = ['Название стартап-проекта', 'Ссылка']
    ws.append(headers)

    for item in data:
        row = [item['Название стартап-проекта'], item['Ссылка']]
        ws.append(row)

    wb.save(excel_filename)


if __name__ == '__main__':
    save_to_excel(process_dir(PATH), OUTPUT_PATH)
